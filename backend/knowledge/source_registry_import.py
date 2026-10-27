from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, asdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from .models import SourceRegistry, ValidationStatus
from .repository import KnowledgeRepository

_REQUIRED = ("source_id", "authority", "official_url", "document_title", "document_type", "jurisdiction")
_ALIASES = {
    "source_id": "source_id",
    "scheme_id": "scheme_id",
    "topic": "topic",
    "scheme": "scheme_name",
    "scheme_name": "scheme_name",
    "authority": "authority",
    "department": "department",
    "official_source_name": "official_source_name",
    "source_name": "official_source_name",
    "official_url": "official_url",
    "url": "official_url",
    "jurisdiction": "jurisdiction",
    "source_type": "source_type",
    "authority_tier": "authority_tier",
    "monitoring_frequency": "monitoring_frequency",
    "monitoring_frequency_days": "monitoring_frequency_days",
    "last_checked": "last_checked",
    "last_changed": "last_changed",
    "current_version_id": "current_version_id",
    "status": "status",
    "document_title": "document_title",
    "document_type": "document_type",
    "document_number": "document_number",
    "publication_date": "publication_date",
    "published_date": "publication_date",
    "effective_date": "effective_from",
    "effective_from": "effective_from",
    "effective_until": "effective_until",
    "dynamic_flag": "dynamic_flag",
    "dynamic": "dynamic_flag",
    "validation_status": "validation_status",
    "rag_eligibility": "rag_eligibility",
    "rag_allowed": "rag_allowed",
    "structured_data_requirement": "structured_data_requirement",
    "structured_db_required": "structured_db_required",
    "replacement_required": "replacement_required",
    "manual_verification": "manual_verification",
    "data_type": "data_type",
    "state": "state",
    "notes": "notes",
}


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), tzinfo=timezone.utc)
    parsed = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _bool(value: Any) -> bool:
    return str(value).strip().lower() not in {"0", "false", "no", "n", "off"}


@dataclass
class ImportReport:
    total_sources: int = 0
    new_sources: int = 0
    updated_sources: int = 0
    skipped_sources: int = 0
    duplicate_sources: int = 0
    missing_urls: int = 0
    missing_document_references: int = 0
    pending_sources: int = 0
    verified_sources: int = 0

    def as_dict(self) -> dict[str, int]:
        return asdict(self)


class SourceRegistryImporter:
    def __init__(self, repository: KnowledgeRepository):
        self.repository = repository

    def import_xlsx(self, workbook_path: str | Path) -> ImportReport:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install dependencies from requirements.txt to import .xlsx files") from exc

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return ImportReport()
        mapped_headers = [_ALIASES.get(_key(header)) for header in headers]
        report = ImportReport()
        seen: set[str] = set()
        for values in rows:
            raw = {field: values[index] for index, field in enumerate(mapped_headers) if field}
            source_id = _text(raw.get("source_id"))
            if not source_id:
                report.skipped_sources += 1
                continue
            report.total_sources += 1
            if source_id in seen:
                report.duplicate_sources += 1
                continue
            seen.add(source_id)
            existing = self.repository.connection.execute(
                "SELECT * FROM source_registry WHERE source_id = ?", (source_id,)
            ).fetchone()
            source = self._source(raw, existing)
            missing_url = not source.official_url.startswith(("http://", "https://"))
            missing_reference = not source.document_title or not source.document_type or not source.document_number
            if missing_url:
                report.missing_urls += 1
            if missing_reference:
                report.missing_document_references += 1

            if existing and existing["validation_status"] == ValidationStatus.VERIFIED.value:
                source = self._source(raw, existing, force_validation=ValidationStatus.VERIFIED.value)
            else:
                source = self._source(raw, existing, force_validation=ValidationStatus.PENDING.value)

            if existing and self._same_registry(existing, source):
                report.skipped_sources += 1
            else:
                self.repository.register_source(source)
                if existing:
                    report.updated_sources += 1
                else:
                    report.new_sources += 1
            if source.validation_status == ValidationStatus.VERIFIED.value:
                report.verified_sources += 1
            else:
                report.pending_sources += 1
        workbook.close()
        return report

    def _source(self, raw: dict[str, Any], existing: Any, *, force_validation: str | None = None) -> SourceRegistry:
        def value(field: str, default: Any = None) -> Any:
            incoming = raw.get(field)
            if incoming not in (None, ""):
                return incoming
            return existing[field] if existing and field in existing.keys() else default

        authority_tier = int(value("authority_tier", 3))
        return SourceRegistry(
            source_id=_text(raw.get("source_id")),
            authority=_text(value("authority", "")),
            source_name=_text(value("official_source_name", value("scheme_name", ""))),
            official_url=_text(value("official_url", "")),
            jurisdiction=_text(value("jurisdiction", "")),
            source_type=_text(value("source_type", "official_document")),
            authority_tier=authority_tier,
            monitoring_frequency_days=int(value("monitoring_frequency_days", 30)),
            last_checked=_date(value("last_checked")),
            last_changed=_date(value("last_changed")),
            current_version_id=_text(value("current_version_id", "")) or None,
            status=_text(value("status", "active")),
            scheme_id=_text(value("scheme_id", "")) or None,
            scheme_name=_text(value("scheme_name", "")) or None,
            official_source_name=_text(value("official_source_name", "")) or None,
            validation_status=force_validation or ValidationStatus.PENDING.value,
            monitoring_frequency=_text(value("monitoring_frequency", "")) or None,
            department=_text(value("department", "")) or None,
            document_title=_text(value("document_title", "")) or None,
            document_type=_text(value("document_type", "")) or None,
            document_number=_text(value("document_number", "")) or None,
            publication_date=_date(value("publication_date")),
            effective_from=_date(value("effective_from")),
            effective_until=_date(value("effective_until")),
            dynamic_flag=_bool(value("dynamic_flag", True)),
            rag_eligibility=_text(value("rag_eligibility", "NOT_ELIGIBLE")) or "NOT_ELIGIBLE",
            structured_data_requirement=_text(value("structured_data_requirement", "")) or None,
            topic=_text(value("topic", "")) or None,
            state=_text(value("state", "")) or None,
            data_type=_text(value("data_type", "")) or None,
            manual_verification=_text(value("manual_verification", "NO")) or "NO",
            rag_allowed=_bool(value("rag_allowed", False)),
            structured_db_required=_bool(value("structured_db_required", False)),
            replacement_required=_text(value("replacement_required", "")) or None,
            notes=_text(value("notes", "")) or None,
        )

    @staticmethod
    def _same_registry(existing: Any, source: SourceRegistry) -> bool:
        fields = {
            "authority": source.authority,
            "source_name": source.source_name,
            "official_url": source.official_url,
            "jurisdiction": source.jurisdiction,
            "source_type": source.source_type,
            "authority_tier": source.authority_tier,
            "monitoring_frequency_days": source.monitoring_frequency_days,
            "current_version_id": source.current_version_id,
            "status": source.status,
            "scheme_id": source.scheme_id,
            "scheme_name": source.scheme_name,
            "official_source_name": source.official_source_name,
            "validation_status": source.validation_status,
            "monitoring_frequency": source.monitoring_frequency,
            "department": source.department,
            "document_title": source.document_title,
            "document_type": source.document_type,
            "document_number": source.document_number,
            "publication_date": source.publication_date.isoformat() if source.publication_date else None,
            "effective_from": source.effective_from.isoformat() if source.effective_from else None,
            "effective_until": source.effective_until.isoformat() if source.effective_until else None,
            "dynamic_flag": int(source.dynamic_flag),
            "rag_eligibility": source.rag_eligibility,
            "structured_data_requirement": source.structured_data_requirement,
            "topic": source.topic,
            "state": source.state,
            "data_type": source.data_type,
            "manual_verification": source.manual_verification,
            "rag_allowed": int(source.rag_allowed),
            "structured_db_required": int(source.structured_db_required),
            "replacement_required": source.replacement_required,
            "notes": source.notes,
        }
        for field, value in fields.items():
            existing_value = existing[field]
            if field in {"publication_date", "effective_from", "effective_until"}:
                existing_value = existing_value
            if existing_value != value:
                return False
        return True


def main() -> None:
    parser = argparse.ArgumentParser(description="Import an authoritative source registry workbook")
    parser.add_argument("workbook")
    parser.add_argument("--db", required=True)
    args = parser.parse_args()
    repository = KnowledgeRepository.from_path(args.db)
    try:
        report = SourceRegistryImporter(repository).import_xlsx(args.workbook)
        print(json.dumps(report.as_dict(), indent=2))
    finally:
        repository.close()


if __name__ == "__main__":
    main()
